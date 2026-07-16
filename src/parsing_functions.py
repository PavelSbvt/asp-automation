import os
import traceback
import re
from typing import Dict, Optional, List, Tuple

import aspose.words as aw
import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication, QMessageBox

from general_functions import _extract_commission_members


def parsing_universal_excel(self) -> None:
    """
    Парсит универсальную таблицу Excel
    :return: None
    """

    try:
        print("[I] НАЧАЛО parsing_universal_excel")

        print(f"[I] Путь к файлу: {self.universal_file_path}")

        # Проверяем существование файла
        if not os.path.exists(self.universal_file_path):
            print(f"[E] Файл не существует: {self.universal_file_path}")
            QMessageBox.critical(self, "Ошибка",
                                 f"Файл не существует:\n{self.universal_file_path}")

            return

        print("[I] Файл существует")

        # Проверяем размер файла
        file_size = os.path.getsize(self.universal_file_path)
        print(f"[I] Размер файла: {file_size} байт")

        if file_size == 0:
            print("[E] Файл пустой!")
            QMessageBox.critical(self, "Ошибка", "Файл пустой!")
            return

        print("[I] Начало чтения файла")

        try:
            print("[I] Через openpyxl...")
            wb = load_workbook(self.universal_file_path, read_only=True, data_only=True)
            ws = wb.active
            print(f"[I] Активный лист: {ws.title}")
            QApplication.processEvents()

            data = []
            for row in ws.iter_rows(values_only=True):
                row_data = []
                for cell in row:
                    row_data.append(str(cell) if cell is not None else "")
                data.append(row_data)
                if len(data) > 100:  # ограничиваем для теста
                    break

            wb.close()
            print(f"[I] Прочитано {len(data)} строк")

            if len(data) > 1:
                df = pd.DataFrame(data[1:], columns=data[0])
            else:
                df = pd.DataFrame()

        except Exception as e:
            print(f"[E] openpyxl не сработал: {e}")

            # Пробуем через pandas
            try:
                print("[I] Через pandas...")
                df = pd.read_excel(
                    self.universal_file_path,
                    dtype=str,
                    engine='openpyxl'
                )
                print(f"[I] pandas прочитал {len(df)} строк")
            except Exception as e2:
                print(f"[E] pandas не сработал: {e2}")
                raise
                # ключевое слово для генерации (выброса) исключений.

        if df is None or len(df) == 0:
            print("[E] Нет данных в файле!")
            QMessageBox.warning(self, "Ошибка", "Файл пуст или не содержит данных!")
            print("[E] Ошибка. Файл пуст или не содержит данных!")
            return

        print(f"[I] Загружено {len(df)} строк, колонок: {len(df.columns)}")

        students_data = []
        for index, row in df.iterrows():
            try:
                if len(row) > 2:
                    fio = str(row.iloc[2]).strip()
                else:
                    fio = ""

                if len(row) > 22:
                    profile = str(row.iloc[22]).strip()
                else:
                    profile = ""

                if len(row) > 7:
                    citizenship = str(row.iloc[7]).strip()
                else:
                    citizenship = ""

                if fio and fio != "nan" and fio != "":
                    if profile == "nan":
                        profile = ""
                    if citizenship == "nan" or citizenship == "":
                        citizenship = "РФ"

                    students_data.append({
                        'fio': fio,
                        'profile': profile,
                        'citizenship': citizenship
                    })
            except Exception as e:
                print(f"[E] Ошибка строки {index}: {e}")
                continue

        self.students_data = students_data
        print(f"[I] Всего студентов: {len(students_data)}")
        QApplication.processEvents()

        for i, student in enumerate(students_data[:5], 1):
            print(f"[I]   {i}. {student['fio']}")

        QMessageBox.information(
            self,
            "Готово!",
            f"Данные успешно загружены!\n\n"
            f"Найдено записей: {len(students_data)}"
        )
        print(f"[S] Данные успешно загружены! Найдено записей: {len(students_data)}")

    except Exception as e:
        print(f"[E] КРИТИЧЕСКАЯ ОШИБКА: {e}")
        traceback.print_exc()
        QApplication.processEvents()
        QMessageBox.critical(
            self,
            "Ошибка",
            f"Не удалось прочитать файл:\n{str(e)}"
        )


def parse_doc_with_tables(path: str) -> Tuple[pd.DataFrame, str]:
    """
    Извлекает все таблицы из документа word .doc и возвращает
    их как один DataFrame с текстом.
    """
    doc = aw.Document(path)
    text = doc.get_text()

    tables_data = []

    # Получаем все таблицы в документе
    for idx, node in enumerate(doc.get_child_nodes(aw.NodeType.TABLE, True)):
        table = node.as_table()
        rows = []

        for row_node in table.rows:
            row = row_node.as_row()
            cells = []
            for cell_node in row.cells:
                cell = cell_node.as_cell()
                # Получаем текст ячейки и очищаем от лишних пробелов
                cell_text = cell.get_text().strip()
                # Убираем управляющие символы
                cell_text = re.sub(r'[\x00-\x1f\x7f]', '', cell_text)
                cells.append(cell_text)
            rows.append(cells)

        if not rows:
            continue

        # Если есть заголовок — делаем его колонками
        if len(rows) > 1 and all(len(r) == len(rows[0]) for r in rows):
            # Проверяем, что первая строка — это заголовок
            header = rows[0]
            # Очищаем заголовки от мусора
            clean_header = []
            for h in header:
                h_clean = re.sub(r'[\x00-\x1f\x7f]', '', h).strip()
                clean_header.append(h_clean)
            df = pd.DataFrame(rows[1:], columns=clean_header)
        else:
            df = pd.DataFrame(rows)

        df['source_file'] = path
        df['table_index'] = idx
        tables_data.append(df)

    tables_df = pd.concat(tables_data, ignore_index=True) if tables_data else pd.DataFrame()
    return tables_df, text


def parsing_commissions(self) -> Optional[Dict[str, Dict]]:
    """
    Парсит файл комиссии и возвращает словарь:
    { "шифр": {"members": [...], "code": "шифр"} }
    """
    print("[I] Парсинг комиссии...")

    # Определяем файл в зависимости от семестра
    if self.radio_button_1.isChecked():
        path = getattr(self, "history_commission_path", None)
        commission_name = "История"
        semester_name = "Первый семестр"
    elif self.radio_button_2.isChecked():
        path = getattr(self, "foreign_commission_path", None)
        commission_name = "Иностранный"
        semester_name = "Второй семестр"
    else:
        path = getattr(self, "specialization_commission_path", None)
        commission_name = "Специальность"
        semester_name = "Третий семестр"

    if not path:
        QMessageBox.warning(self, "Ошибка",
                            f"Сначала выберите файл комиссии '{commission_name}'!")
        print("[E] Ошибка. сначала выберите файл комиссии")
        return None

    if not os.path.exists(path):
        QMessageBox.warning(self, "Ошибка", f"Файл не существует: {path}")
        print("[E] ошибка. файл не существует")
        return None

    try:
        tables_df, full_text = parse_doc_with_tables(path)
        commissions_by_code = {}

        # ===== ОТЛАДКА =====
        print("[I] ===== СОДЕРЖИМОЕ ТАБЛИЦЫ =====")
        for idx, row in tables_df.iterrows():
            print(f"[I] Строка {idx}:")
            for col in tables_df.columns:
                val = str(row[col])[:100] if pd.notna(row[col]) else ""
                print(f"    {col}: {val}")
            print("-" * 40)
        # ===== КОНЕЦ ОТЛАДКИ =====

        if tables_df is not None and not tables_df.empty:
            cols = list(tables_df.columns)
            print(f"[I] Колонки: {cols}")

            # Находим колонки
            fio_col = None
            code_col = None

            for col in cols:
                if col not in ['source_file', 'table_index']:
                    col_lower = col.lower()
                    if 'фамил' in col_lower or 'инициал' in col_lower:
                        fio_col = col
                    if 'шифр' in col_lower or 'наименование' in col_lower or 'специальност' in col_lower:
                        code_col = col

            # Если не нашли — берём по индексам
            if code_col is None and len(cols) >= 3:
                code_col = cols[1]
            if fio_col is None and len(cols) >= 2:
                fio_col = cols[0]

            print(f"[I] Колонка ФИО: {fio_col}")
            print(f"[I] Колонка шифра: {code_col}")

            # Проходим по строкам
            for idx, row in tables_df.iterrows():
                if fio_col is None:
                    for col in cols:
                        if col not in ['source_file', 'table_index']:
                            value = str(row[col]).strip() if pd.notna(row[col]) else ""
                            if value and len(value) > 10:
                                _extract_commission_members(value, commissions_by_code, row, code_col)
                else:
                    value = str(row[fio_col]).strip() if pd.notna(row[fio_col]) else ""
                    if value and len(value) > 10:
                        _extract_commission_members(value, commissions_by_code, row, code_col)

        self.commissions_by_code = commissions_by_code

        if commissions_by_code:
            print(f"[I] Найдено комиссий для {len(commissions_by_code)} шифров:")
            for code, comm in commissions_by_code.items():
                members = comm.get('members', [])
                print(f"[I]   Шифр: {code} -> {len(members)} членов")
                for member in members:
                    print(f"[I]     - {member.get('fio', '')} ({member.get('role', '')})")

            print(f"[S] Готово. Комиссия успешно загружена. Найдено комиссий: "
                  f"{len(commissions_by_code)}. Семестр: {semester_name}.")
        else:
            print("[E] Члены комиссии не найдены!")

        return commissions_by_code

    except Exception as e:
        print(f"[E] Ошибка при парсинге комиссии: {e}")
        traceback.print_exc()
        QMessageBox.critical(
            self,
            "Ошибка",
            f"Не удалось прочитать файл комиссии:\n{str(e)}"
        )
        return None

def site_parce_dn() -> List[Tuple[str, str]]:
    """
    парсинг данных профилей и их кодов аспирантов с сайта магистратуры
    :return: :return: List[Tuple[str, str]]
    """

    url = 'http://dn.tstu.ru/index.php/specialties'
    response = requests.get(url, timeout=10)
    response.raise_for_status()

    soup = BeautifulSoup(response.text, 'html.parser')
    target_table = None
    for table in soup.find_all('table'): #Цикл по всем таблицам на сайте
        headers = [th.get_text(strip=True) for th in table.find_all('th')]
        if 'Шифр и наименование научной специальности' in headers: #Поиск нужной таблицы
            target_table = table
            break
    if target_table is None:
        return []

    specialties = []
    for row in target_table.find_all('tr'): #Добавление кода, профиль в список как кортеж
        cols = row.find_all('td')
        if not cols:
            continue

        #Проверяем сколько яйчеек в строке
        if len(cols) == 3: #Если 3, то берём код и специальность из 2 и 3
            code = cols[1].get_text(strip=True)
            name = cols[2].get_text(strip=True)
            if code and code.endswith('.') and ' ' not in code:
                specialties.append((code, name))
        elif len(cols) == 2: #Если 2, то код и специальность из 1 и 2
            code = cols[0].get_text(strip=True)
            name = cols[1].get_text(strip=True)
            if code and code.endswith('.') and ' ' not in code:
                specialties.append((code, name))
    return specialties

def site_parce_main() -> List[Tuple[str, str]]:
    """
    парсинг данных (аспирантов) профилей и их кодов с сайта ТГТУ
    :return: List[Tuple[str, str]]
    """

    url = 'https://www.tstu.ru/r.php?r=postup.abitur.na_sp'
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,'
                'application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7',
    } #Имитируем браузер
    response = requests.get(url, headers=headers, timeout=10)
    #Пытаемся получить ответ от сайта
    response.raise_for_status()

    soup = BeautifulSoup(response.text, 'html.parser')

    target_table = None
    for header in soup.find_all(['h4', 'font'], class_='head4'):
        # цикл по всем элементам, чтобы найти таблицу с заголовком "Аспирантура"
        if header.get_text(strip=True) == 'Аспирантура':
            target_table = header.find_next('table')
            break

    if target_table is None: #Поиск по якорю, если не нашлось по тегам
        anchor = soup.find('a', attrs={'name': '4'})
        if anchor:
            next_tag = anchor.find_next()
            while next_tag and next_tag.name != 'table':
                next_tag = next_tag.find_next()
            target_table = next_tag

    specialties = []
    for row in target_table.find_all('tr'):
        # Добавляем каждый код и каждый профиль в список как кортеж
        cols = row.find_all('td')
        if not cols:
            continue
        if len(cols) >= 2: #Проверяем сколько яйчеек в строке
            code = cols[-2].get_text(strip=True) #Код профиля из предпоследней ячейки
            name = cols[-1].get_text(strip=True) #Название профиля из последней ячейки
            if code and '.' in code and ' ' not in code:
                specialties.append((code, name))

    return specialties
